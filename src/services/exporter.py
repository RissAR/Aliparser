"""Excel (.xlsx) export service for scraped AliExpress products."""

import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.db.models import Product

EXCEL_HEADERS = [
    "ID товару",
    "Назва товару",
    "Ціна",
    "Стара ціна",
    "Валюта",
    "Рейтинг",
    "Замовлення",
    "Магазин",
    "Категорія / Запит",
    "Посилання на товар",
    "Посилання на фото",
    "Дата збору (UTC)",
]


def generate_products_excel(products: List[Product]) -> bytes:
    """Generate a formatted .xlsx workbook with clickable product links.

    Args:
        products: List of Product ORM objects.

    Returns:
        Raw bytes of the generated .xlsx file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "AliExpress Products"

    # Freeze header row
    ws.freeze_panes = "A2"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.append(EXCEL_HEADERS)
    for col_idx in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    link_font = Font(color="2563EB", underline="single")

    for row_idx, product in enumerate(products, start=2):
        scraped_str = (
            product.scraped_at.strftime("%Y-%m-%d %H:%M:%S")
            if product.scraped_at is not None
            else ""
        )

        row_values = [
            product.item_id,
            product.title,
            float(product.price) if product.price is not None else None,
            float(product.original_price) if product.original_price is not None else None,
            product.currency or "USD",
            float(product.rating) if product.rating is not None else None,
            product.orders_count or "",
            product.store_name or "",
            product.category_or_query or "",
            product.product_url or "",
            product.image_url or "",
            scraped_str,
        ]
        ws.append(row_values)

        # Make product URL column (column 10) a clickable hyperlink
        if product.product_url:
            url_cell = ws.cell(row=row_idx, column=10)
            url_cell.hyperlink = product.product_url
            url_cell.font = link_font

    # Auto-adjust column widths within [12, 50]
    min_width = 12
    max_width = 50
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value) if cell.value is not None else ""
            if len(val_str) > max_len:
                max_len = len(val_str)
        adjusted_width = max(min_width, min(max_width, max_len + 3))
        ws.column_dimensions[col_letter].width = adjusted_width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
