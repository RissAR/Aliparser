"""Comprehensive verification script for FastAPI REST endpoints and Excel (.xlsx) export service."""

import io
import logging
import os
import sys

# Ensure project root is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete

from src.api.app import app
from src.db.database import SessionLocal, init_db, upsert_products
from src.db.models import ParsingTask, Product, Proxy
from src.parser.aliexpress import AliProductItem

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("test_api_and_export")


def run_tests() -> None:
    """Execute end-to-end verification of REST endpoints and Excel export."""
    init_db()
    client = TestClient(app)

    # Clean up any leftover test data
    with SessionLocal() as db:
        db.execute(delete(Product).where(Product.item_id.in_(["API_TEST_101", "API_TEST_102"])))
        db.execute(delete(Proxy).where(Proxy.raw_proxy == "10.20.30.40:8080:myuser:mypass"))
        db.commit()

    # 1. Test GET /api/parser/status & GET /api/parser/tasks
    logger.info("Step 1: Testing GET /api/parser/status and GET /api/parser/tasks...")
    res_status = client.get("/api/parser/status")
    assert res_status.status_code == 200, f"Expected 200, got {res_status.status_code}"
    status_data = res_status.json()
    for field in (
        "is_running",
        "status",
        "current_page",
        "max_pages",
        "completed_pages",
        "progress_percent",
        "items_scraped",
        "elapsed_seconds",
        "items_per_second",
    ):
        assert field in status_data, f"Missing '{field}' in status response"
    logger.info("  [OK] GET /api/parser/status returned valid schema: status=%s", status_data["status"])

    res_tasks = client.get("/api/parser/tasks")
    assert res_tasks.status_code == 200
    assert isinstance(res_tasks.json(), list)
    logger.info("  [OK] GET /api/parser/tasks returned %d records", len(res_tasks.json()))

    # 2. Test Proxy Management endpoints
    logger.info("Step 2: Testing Proxy CRUD endpoints (/api/proxies)...")
    res_proxy_add = client.post(
        "/api/proxies",
        json={"raw_proxy": "10.20.30.40:8080:myuser:mypass", "check_connectivity": False},
    )
    assert res_proxy_add.status_code == 200, f"Expected 200, got {res_proxy_add.status_code}: {res_proxy_add.text}"
    proxy_json = res_proxy_add.json()
    proxy_id = proxy_json["id"]
    assert proxy_json["host"] == "10.20.30.40"
    assert proxy_json["port"] == 8080
    assert proxy_json["username"] == "myuser"
    assert proxy_json["protocol"] == "http"
    logger.info("  [OK] POST /api/proxies created proxy ID=%d", proxy_id)

    # Invalid proxy should return 400 Bad Request
    res_proxy_bad = client.post("/api/proxies", json={"raw_proxy": "invalid_proxy_str"})
    assert res_proxy_bad.status_code == 400, f"Expected 400 for invalid proxy, got {res_proxy_bad.status_code}"
    logger.info("  [OK] POST /api/proxies rejected invalid proxy with 400 Bad Request")

    # List proxies
    res_proxy_list = client.get("/api/proxies")
    assert res_proxy_list.status_code == 200
    proxies_list = res_proxy_list.json()
    assert any(p["id"] == proxy_id for p in proxies_list), "Created proxy not found in GET /api/proxies"
    logger.info("  [OK] GET /api/proxies returned %d proxies (including ID=%d)", len(proxies_list), proxy_id)

    # Delete proxy
    res_proxy_del = client.delete(f"/api/proxies/{proxy_id}")
    assert res_proxy_del.status_code == 200
    assert res_proxy_del.json() == {"deleted": True, "id": proxy_id}
    logger.info("  [OK] DELETE /api/proxies/%d succeeded", proxy_id)

    # 3. Test Products insertion, listing, and search filtering
    logger.info("Step 3: Testing GET /api/products and search filtering...")
    with SessionLocal() as db:
        test_task = ParsingTask(target="gaming mouse", max_pages=1, current_page=1, status="completed", items_scraped=2)
        db.add(test_task)
        db.commit()
        db.refresh(test_task)
        test_task_id = test_task.id

        sample_items = [
            AliProductItem(
                item_id="API_TEST_101",
                title="Logitech G Pro X Superlight Wireless Gaming Mouse",
                price=119.99,
                original_price=159.99,
                currency="USD",
                rating=4.9,
                orders_count="5,000+ sold",
                store_name="ProGear Official Store",
                product_url="https://www.aliexpress.com/item/API_TEST_101.html",
                image_url="https://ae01.alicdn.com/kf/api101.jpg",
            ),
            AliProductItem(
                item_id="API_TEST_102",
                title="Baseus 100W GaN Fast Charger Type-C Adapter",
                price=29.50,
                original_price=45.00,
                currency="USD",
                rating=4.8,
                orders_count="10,000+ sold",
                store_name="Baseus Official Store",
                product_url="https://www.aliexpress.com/item/API_TEST_102.html",
                image_url="https://ae01.alicdn.com/kf/api102.jpg",
            ),
        ]
        upsert_products(db, sample_items, category_or_query="gaming accessories", task_id=test_task_id)

    res_products = client.get(f"/api/products?task_id={test_task_id}")
    assert res_products.status_code == 200
    prod_data = res_products.json()
    assert prod_data["total"] == 2
    assert len(prod_data["items"]) == 2
    logger.info("  [OK] GET /api/products?task_id=%d returned total=%d", test_task_id, prod_data["total"])

    res_search = client.get(f"/api/products?search=Superlight&task_id={test_task_id}")
    assert res_search.status_code == 200
    search_data = res_search.json()
    assert search_data["total"] == 1
    assert search_data["items"][0]["item_id"] == "API_TEST_101"
    logger.info("  [OK] GET /api/products?search=Superlight matched 1 item ('%s')", search_data["items"][0]["title"])

    # 4. Test Excel (.xlsx) Export
    logger.info("Step 4: Testing GET /api/products/export and inspecting .xlsx structure...")
    res_export = client.get(f"/api/products/export?task_id={test_task_id}")
    assert res_export.status_code == 200
    assert (
        res_export.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "aliexpress_products.xlsx" in res_export.headers.get("content-disposition", "")

    wb = load_workbook(io.BytesIO(res_export.content))
    assert "AliExpress Products" in wb.sheetnames, f"Missing sheet 'AliExpress Products': {wb.sheetnames}"
    ws = wb["AliExpress Products"]

    assert ws.freeze_panes == "A2", f"Expected freeze_panes='A2', got {ws.freeze_panes}"

    headers = [ws.cell(row=1, column=col).value for col in range(1, 13)]
    expected_headers = [
        "Item ID",
        "Product Title",
        "Price",
        "Original Price",
        "Currency",
        "Rating",
        "Orders",
        "Store",
        "Category / Query",
        "Product URL",
        "Image URL",
        "Scraped At (UTC)",
    ]
    assert headers == expected_headers, f"Excel headers mismatch: {headers} != {expected_headers}"
    logger.info("  [OK] Excel sheet 'AliExpress Products' headers verified (12 English columns)")

    row2_id = ws.cell(row=2, column=1).value
    row2_title = ws.cell(row=2, column=2).value
    row2_price = ws.cell(row=2, column=3).value
    row2_url_cell = ws.cell(row=2, column=10)

    assert row2_id in ("API_TEST_101", "API_TEST_102")
    assert isinstance(row2_price, (int, float)) and row2_price > 0
    assert row2_url_cell.hyperlink is not None, "Product URL cell in column 10 is missing hyperlink!"
    assert row2_url_cell.hyperlink.target.startswith("https://www.aliexpress.com/item/API_TEST_")
    logger.info(
        "  [OK] Row 2 verified: ID=%s, Title='%s', Price=%s, Clickable Hyperlink='%s'",
        row2_id,
        row2_title[:30],
        row2_price,
        row2_url_cell.hyperlink.target,
    )

    # Clean up test products and task via API and DB
    res_del_prods = client.delete(f"/api/products?task_id={test_task_id}")
    assert res_del_prods.status_code == 200
    assert res_del_prods.json()["deleted"] == 2
    with SessionLocal() as db:
        db.execute(delete(ParsingTask).where(ParsingTask.id == test_task_id))
        db.commit()
    logger.info("  [OK] Cleaned up temporary test products and task.")

    print("\nALL API & EXCEL EXPORT TESTS PASSED")


if __name__ == "__main__":
    run_tests()
