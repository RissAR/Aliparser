"""Comprehensive End-to-End (E2E) and architectural audit test suite for AliParser."""

import asyncio
import io
import logging
import os
import sys
import zipfile

# Ensure project root is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import httpx
from httpx import ASGITransport
from openpyxl import load_workbook
from sqlalchemy import delete, select, text

from src.api.app import app
from src.db.database import BADGE_IMAGE_SUBSTRINGS, SessionLocal, init_db
from src.db.models import ParsingTask, Product, Proxy
from src.services.exporter import EXCEL_HEADERS

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("test_e2e_full")


async def run_e2e_audit() -> None:
    """Execute full E2E and architectural verification across all 5 audit stages."""
    init_db()

    # =========================================================================
    # Stage 1: Verify SQLite WAL Mode & Stale Task Recovery
    # =========================================================================
    logger.info("Stage 1: Verifying SQLite WAL mode and stale task recovery...")
    with SessionLocal() as db:
        journal_mode = db.execute(text("PRAGMA journal_mode;")).scalar_one()
        assert str(journal_mode).lower() == "wal", f"Expected journal_mode='wal', got '{journal_mode}'"
        logger.info("  [OK] SQLite PRAGMA journal_mode = '%s'", journal_mode)

        # Insert a simulated stale running task and verify init_db() recovers it to 'stopped'
        stale_task = ParsingTask(target="stale_recovery_test", max_pages=2, current_page=1, status="running")
        db.add(stale_task)
        db.commit()
        db.refresh(stale_task)
        stale_id = stale_task.id

    init_db()

    with SessionLocal() as db:
        recovered = db.execute(select(ParsingTask).where(ParsingTask.id == stale_id)).scalar_one()
        assert recovered.status == "stopped", f"Expected stale task status='stopped', got '{recovered.status}'"
        assert recovered.finished_at is not None, "Expected finished_at to be populated on recovered stale task"
        db.execute(delete(ParsingTask).where(ParsingTask.id == stale_id))
        db.commit()
        logger.info("  [OK] Stale 'running' task automatically recovered to 'stopped' on init_db()")

    transport = ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        # =====================================================================
        # Stage 2: E2E Real Search Query Scrape via FastAPI + TaskRunner
        # =====================================================================
        logger.info("Stage 2: E2E Real Search Query Scrape ('wireless mouse', max_pages=1)...")
        res_start = await client.post(
            "/api/parser/start",
            json={"target": "wireless mouse", "max_pages": 1, "delay_seconds": 0.5},
        )
        assert res_start.status_code == 200, f"Expected 200 on start, got {res_start.status_code}: {res_start.text}"
        start_status = res_start.json()
        task1_id = start_status["task_id"]
        assert task1_id is not None

        final_status1 = start_status
        for _ in range(180):
            await asyncio.sleep(0.5)
            res_stat = await client.get("/api/parser/status")
            assert res_stat.status_code == 200
            final_status1 = res_stat.json()
            assert "is_running" in final_status1
            assert "elapsed_seconds" in final_status1
            assert "estimated_remaining_seconds" in final_status1
            if not final_status1["is_running"]:
                break

        assert final_status1["status"] == "completed", f"Expected 'completed', got: {final_status1}"
        assert final_status1["items_scraped"] >= 15, (
            f"Expected >= 15 items scraped for 'wireless mouse', got {final_status1['items_scraped']}"
        )
        logger.info(
            "  [OK] Task #%d completed: scraped %d real items in %.2fs",
            task1_id,
            final_status1["items_scraped"],
            final_status1["elapsed_seconds"],
        )

        # Verify no promotional badge images exist in scraped items
        res_prods1 = await client.get(f"/api/products?task_id={task1_id}&limit=200")
        assert res_prods1.status_code == 200
        prods1 = res_prods1.json()["items"]
        assert len(prods1) >= 15
        for item in prods1:
            img_url = item.get("image_url") or ""
            for forbidden in ("154x64", "48x48", "S1887a285b60743859ac7bdbfca5e0896Z"):
                assert forbidden not in img_url, f"Found forbidden badge pattern '{forbidden}' in image_url: {img_url}"
            for badge_pat in BADGE_IMAGE_SUBSTRINGS:
                assert badge_pat not in img_url, f"Found badge substring '{badge_pat}' in image_url: {img_url}"
        logger.info("  [OK] Verified %d scraped products have clean image_url values (0 badges)", len(prods1))

        # =====================================================================
        # Stage 3: E2E Real Category/Catalog Direct URL Scrape
        # =====================================================================
        direct_url = "https://www.aliexpress.com/w/wholesale-usb-flash-drive.html"
        logger.info("Stage 3: E2E Direct URL Scrape ('%s', max_pages=1)...", direct_url)
        res_start2 = await client.post(
            "/api/parser/start",
            json={"target": direct_url, "max_pages": 1, "delay_seconds": 0.5},
        )
        assert res_start2.status_code == 200, f"Expected 200 on start, got {res_start2.status_code}: {res_start2.text}"
        task2_id = res_start2.json()["task_id"]

        final_status2 = res_start2.json()
        for _ in range(180):
            await asyncio.sleep(0.5)
            res_stat2 = await client.get("/api/parser/status")
            assert res_stat2.status_code == 200
            final_status2 = res_stat2.json()
            if not final_status2["is_running"]:
                break

        assert final_status2["status"] == "completed", f"Expected 'completed', got: {final_status2}"
        assert final_status2["items_scraped"] >= 15, (
            f"Expected >= 15 items scraped for direct URL, got {final_status2['items_scraped']}"
        )
        logger.info(
            "  [OK] Direct URL Task #%d completed: scraped %d real items in %.2fs",
            task2_id,
            final_status2["items_scraped"],
            final_status2["elapsed_seconds"],
        )

        # =====================================================================
        # Stage 4: E2E Search Filtering, Proxy CRUD, and Excel (.xlsx) Export
        # =====================================================================
        logger.info("Stage 4: Verifying product filtering, Proxy CRUD, and real Excel (.xlsx) export...")
        res_search = await client.get("/api/products?search=usb")
        assert res_search.status_code == 200
        search_payload = res_search.json()
        assert search_payload["total"] >= 1, "Expected at least 1 product matching search='usb'"
        logger.info("  [OK] GET /api/products?search=usb matched %d products", search_payload["total"])

        res_by_task = await client.get(f"/api/products?task_id={task2_id}")
        assert res_by_task.status_code == 200
        assert res_by_task.json()["total"] >= 15
        logger.info("  [OK] GET /api/products?task_id=%d returned %d products", task2_id, res_by_task.json()["total"])

        # Proxy CRUD check
        res_proxy = await client.post(
            "/api/proxies",
            json={"raw_proxy": "192.0.2.55:8080:e2euser:e2epass", "check_connectivity": False},
        )
        assert res_proxy.status_code == 200
        proxy_id = res_proxy.json()["id"]
        res_proxy_del = await client.delete(f"/api/proxies/{proxy_id}")
        assert res_proxy_del.status_code == 200
        logger.info("  [OK] Proxy CRUD verified (created & deleted proxy ID=%d)", proxy_id)

        # Excel export check on real scraped data
        res_export = await client.get(f"/api/products/export?task_id={task2_id}")
        assert res_export.status_code == 200
        assert (
            res_export.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        xlsx_bytes = res_export.content

        # Verify normalized relative paths in xl/_rels/workbook.xml.rels
        with zipfile.ZipFile(io.BytesIO(xlsx_bytes), "r") as zf:
            rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            assert 'Target="/xl/' not in rels_xml, "Found non-relative Target=\"/xl/ in workbook.xml.rels"

        wb = load_workbook(io.BytesIO(xlsx_bytes))
        assert "AliExpress Products" in wb.sheetnames
        ws = wb["AliExpress Products"]
        headers = [ws.cell(row=1, column=col).value for col in range(1, 13)]
        assert headers == EXCEL_HEADERS, f"Headers mismatch: {headers} != {EXCEL_HEADERS}"
        assert ws.max_row >= 16, f"Expected at least 16 rows in exported workbook, got {ws.max_row}"

        row2_price = ws.cell(row=2, column=3).value
        row2_url_cell = ws.cell(row=2, column=10)
        assert isinstance(row2_price, (int, float)) and row2_price > 0
        assert row2_url_cell.hyperlink is not None
        assert row2_url_cell.hyperlink.target.startswith("https://www.aliexpress.com/item/")
        logger.info(
            "  [OK] Exported .xlsx verified: %d rows, English headers, numeric prices, clickable links, relative ZIP rels",
            ws.max_row - 1,
        )

        # =====================================================================
        # Stage 5: E2E State Stress Test (409 Conflict on Double Start + Graceful Stop)
        # =====================================================================
        logger.info("Stage 5: Testing 409 Conflict on concurrent start and Graceful Stop...")
        res_start3 = await client.post(
            "/api/parser/start",
            json={"target": "mechanical keyboard", "max_pages": 3, "delay_seconds": 1.5},
        )
        assert res_start3.status_code == 200, f"Expected 200 on start, got {res_start3.status_code}"

        # Immediate second start must return 409 Conflict
        res_conflict = await client.post(
            "/api/parser/start",
            json={"target": " duplicate start", "max_pages": 1, "delay_seconds": 1.0},
        )
        assert res_conflict.status_code == 409, f"Expected 409 Conflict, got {res_conflict.status_code}"
        logger.info("  [OK] Concurrent POST /api/parser/start rejected with 409 Conflict")

        # Trigger Graceful Stop
        await asyncio.sleep(0.5)
        res_stop = await client.post("/api/parser/stop")
        assert res_stop.status_code == 200
        stop_status = res_stop.json()
        assert stop_status["status"] == "stopped", f"Expected status='stopped', got '{stop_status['status']}'"
        assert stop_status["is_running"] is False
        logger.info("  [OK] POST /api/parser/stop transitioned task cleanly to 'stopped'")

    print("\nALL E2E & ARCHITECTURE AUDIT TESTS PASSED")


if __name__ == "__main__":
    asyncio.run(run_e2e_audit())
